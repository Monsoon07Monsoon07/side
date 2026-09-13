import openpyxl, re
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
from copy import copy

SRC = "Excel加工_実戦サンプル_加工前.xlsx"
DST = "Excel加工_実戦サンプル_加工後.xlsx"

wb = openpyxl.load_workbook(SRC, data_only=True)
ws_src = wb["加工前_顧客売上データ"]
header = [c.value for c in ws_src[1]]
raw_rows = list(ws_src.iter_rows(min_row=2, values_only=True))

report = {
    "加工前行数": len(raw_rows),
    "エラー": [],
}

# ---- 1. exact duplicate removal (keep first occurrence, preserve order) ----
seen = set()
deduped = []
dup_removed = 0
for r in raw_rows:
    if r in seen:
        dup_removed += 1
        continue
    seen.add(r)
    deduped.append(list(r))
report["削除した重複行数"] = dup_removed
report["加工後行数(重複削除後)"] = len(deduped)

# ---- name surname dictionary (built ONLY from rows that already contain an explicit space) ----
import collections
surname_counter = collections.Counter()
for r in deduped:
    name = r[2]
    if (" " in name) or ("　" in name):
        parts = [p for p in re.split(r"[ 　]+", name.strip()) if p]
        if len(parts) == 2:
            surname_counter[parts[0]] += 1
known_surnames = sorted(surname_counter.keys(), key=len, reverse=True)

def split_name(name):
    """Return (formatted_name, needs_review: bool)."""
    raw = name
    if (" " in raw) or ("　" in raw):
        parts = [p for p in re.split(r"[ 　]+", raw.strip()) if p]
        if len(parts) == 2:
            return f"{parts[0]} {parts[1]}", False
        return raw, True  # more than 2 parts -> can't confidently reformat
    # no space at all: try to match against surnames CONFIRMED elsewhere in this dataset
    cands = [s for s in known_surnames if raw.startswith(s)]
    if not cands:
        return raw, True
    maxlen = max(len(c) for c in cands)
    longest = [c for c in cands if len(c) == maxlen]
    if len(longest) != 1:
        return raw, True  # genuine ambiguity between equally-long candidate surnames
    sur = longest[0]
    given = raw[len(sur):]
    if not given:
        return raw, True
    return f"{sur} {given}", False

# ---- date normalization ----
REIWA_OFFSET = 2018  # Reiwa 1 = 2019

def normalize_date(v):
    s = str(v).strip()
    m = re.match(r"^(\d{4})-(\d{2})-(\d{2})$", s)
    if m:
        return f"{m.group(1)}-{m.group(2)}-{m.group(3)}", False
    m = re.match(r"^(\d{4})/(\d{1,2})/(\d{1,2})$", s)
    if m:
        return f"{int(m.group(1)):04d}-{int(m.group(2)):02d}-{int(m.group(3)):02d}", False
    m = re.match(r"^令和(\d+)年(\d{1,2})月(\d{1,2})日$", s)
    if m:
        year = REIWA_OFFSET + int(m.group(1))
        return f"{year:04d}-{int(m.group(2)):02d}-{int(m.group(3)):02d}", False
    return s, True  # unrecognized format -> flagged, not guessed

# ---- phone normalization ----
ZEN2HAN_DIGITS = str.maketrans("０１２３４５６７８９－", "0123456789-")

def normalize_phone(v):
    s = str(v).strip()
    s_half = s.translate(ZEN2HAN_DIGITS)
    if "-" in s_half:
        # already explicitly segmented by the source (full-width or half-width hyphen) -> just normalize width, keep the source's own segmentation as-is
        return s_half, False
    digits = re.sub(r"\D", "", s_half)
    if len(digits) == 11:
        # No explicit separator in source. This dataset's OTHER 548 rows are hyphenated
        # and are, without exception, in a 3-4-4 pattern for every 11-digit number
        # regardless of prefix. We apply that dataset-confirmed convention here rather
        # than guessing a real-world area-code boundary.
        return f"{digits[0:3]}-{digits[3:7]}-{digits[7:11]}", False
    return s, True  # not 11 digits after cleanup -> cannot confidently format, flag it

# ---- apply all transforms ----
name_fixed = 0
phone_fixed = 0
date_fixed = 0
product_trimmed = 0
amount_mismatch = 0
review_flags = 0

CHECK_COL_INDEX = len(header)  # append as new last column

processed_rows = []
for r in deduped:
    order_id, order_date, cust_name, phone, email, pref, category, product, qty, price, total, note = r

    new_date, date_ng = normalize_date(order_date)
    if new_date != str(order_date).strip():
        date_fixed += 1

    new_name, name_ng = split_name(cust_name)
    if new_name != cust_name:
        name_fixed += 1

    new_phone, phone_ng = normalize_phone(phone)
    if new_phone != str(phone).strip():
        phone_fixed += 1

    new_product = product.strip()
    if new_product != product:
        product_trimmed += 1

    calc_ok = (qty is not None and price is not None and total is not None and qty * price == total)
    if not calc_ok:
        amount_mismatch += 1
    check_result = "OK" if calc_ok else "要確認"

    if date_ng:
        new_date = f"{order_date}（要確認）"
        review_flags += 1
    if name_ng:
        new_name = f"{cust_name}（要確認）"
        review_flags += 1
    if phone_ng:
        new_phone = f"{phone}（要確認）"
        review_flags += 1

    processed_rows.append([
        order_id, new_date, new_name, new_phone, email, pref, category,
        new_product, qty, price, total, note, check_result
    ])

report["日付変換件数"] = date_fixed
report["顧客名修正件数"] = name_fixed
report["電話番号修正件数"] = phone_fixed
report["商品名空白修正件数"] = product_trimmed
report["売上金額不一致件数"] = amount_mismatch
report["要確認件数(名前/日付/電話の書式判断不能)"] = review_flags
report["加工後行数(最終)"] = len(processed_rows)

# ---- write new sheet, leaving original sheets untouched ----
new_header = header + ["チェック結果"]
if "加工後_顧客売上データ" in wb.sheetnames:
    del wb["加工後_顧客売上データ"]
ws_new = wb.create_sheet("加工後_顧客売上データ")

bold = Font(name="Arial", bold=True, color="FFFFFF")
fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
normal_font = Font(name="Arial")
ng_font = Font(name="Arial", color="C00000")

ws_new.append(new_header)
for cell in ws_new[1]:
    cell.font = bold
    cell.fill = fill
    cell.alignment = Alignment(horizontal="center", vertical="center")

for row in processed_rows:
    ws_new.append(row)

last_row = ws_new.max_row
last_col = ws_new.max_column

for r_idx in range(2, last_row + 1):
    for c_idx in range(1, last_col + 1):
        cell = ws_new.cell(row=r_idx, column=c_idx)
        cell.font = normal_font
        val = cell.value
        if isinstance(val, str) and "要確認" in val:
            cell.font = ng_font
    check_cell = ws_new.cell(row=r_idx, column=last_col)
    if check_cell.value == "要確認":
        check_cell.font = ng_font
    else:
        check_cell.font = Font(name="Arial", color="006100")

# number formats
qty_col = header.index("数量") + 1
price_col = header.index("単価") + 1
total_col = header.index("売上金額") + 1
for r_idx in range(2, last_row + 1):
    ws_new.cell(row=r_idx, column=qty_col).number_format = "0"
    ws_new.cell(row=r_idx, column=price_col).number_format = "#,##0"
    ws_new.cell(row=r_idx, column=total_col).number_format = "#,##0"

# column widths (rough content-based sizing)
widths = {
    "注文ID": 12, "注文日": 13, "顧客名": 16, "電話番号": 16, "メールアドレス": 22,
    "都道府県": 10, "商品カテゴリ": 12, "商品名": 16, "数量": 7, "単価": 10,
    "売上金額": 12, "備考": 14, "チェック結果": 12,
}
for i, h in enumerate(new_header, start=1):
    ws_new.column_dimensions[get_column_letter(i)].width = widths.get(h, 14)

ws_new.freeze_panes = "A2"
ws_new.auto_filter.ref = f"A1:{get_column_letter(last_col)}{last_row}"
ws_new.sheet_view.showGridLines = True

wb.save(DST)
print("SAVED:", DST)

import json
print(json.dumps(report, ensure_ascii=False, indent=2))
